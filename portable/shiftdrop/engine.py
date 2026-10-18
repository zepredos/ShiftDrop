"""Local, collision-safe conversions. No GUI or network access is required."""

from __future__ import annotations

import io
import math
import shutil
import subprocess
import tempfile
from contextlib import closing
from pathlib import Path
from xml.sax.saxutils import escape

from PIL import Image, ImageDraw, ImageEnhance, ImageFilter, ImageOps
from pillow_heif import register_heif_opener
import pypdfium2 as pdfium
from docx import Document
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from striprtf.striprtf import rtf_to_text

register_heif_opener()
Image.init()
MAX_PIXELS = 100_000_000
Image.MAX_IMAGE_PIXELS = MAX_PIXELS
TEXT_EXTENSIONS = {".txt", ".md", ".csv", ".json", ".log", ".rtf", ".docx", ".yaml", ".yml"}
ENCODERS = {"JPG": "JPEG", "PNG": "PNG", "WEBP": "WEBP", "HEIC": "HEIF", "TIFF": "TIFF", "AVIF": "AVIF", "BMP": "BMP"}
EXTENSIONS = {"JPG": "jpg", "PNG": "png", "WEBP": "webp", "HEIC": "heic", "TIFF": "tiff", "AVIF": "avif", "BMP": "bmp", "PDF": "pdf", "DOCX": "docx", "TXT": "txt"}


class ConversionError(ValueError):
    pass


def detect_kind(source: str | Path) -> str:
    path = Path(source)
    if not path.is_file():
        raise ConversionError("Choose an existing file, rather than a folder.")
    if path.suffix.lower() == ".pdf":
        return "pdf"
    if path.suffix.lower() in TEXT_EXTENSIONS:
        return "text"
    if path.suffix.lower() in Image.registered_extensions():
        return "image"
    raise ConversionError("This version supports images, PDFs, and text documents.")


def check_size(width: int, height: int):
    if width < 1 or height < 1 or width > 16000 or height > 16000 or width * height > MAX_PIXELS:
        raise ConversionError("Use an image below 100 megapixels and 16,000 pixels per side.")


def load_image(source: str | Path) -> Image.Image:
    with Image.open(source) as image:
        check_size(*image.size)
        result = ImageOps.exif_transpose(image)
        result.load()
        return result.convert("RGBA")


def encode(image: Image.Image, format_name: str, quality: float = .9) -> bytes:
    check_size(*image.size)
    image = image.copy()
    image.info.clear()
    codec = ENCODERS.get(format_name)
    if codec not in Image.SAVE:
        raise ConversionError(f"{format_name} encoding is unavailable in this installation.")
    if codec in {"JPEG", "BMP"}:
        rgba = image.convert("RGBA")
        image = Image.new("RGB", image.size, "white")
        image.paste(rgba, mask=rgba.getchannel("A"))
    buffer = io.BytesIO()
    image.save(buffer, format=codec, quality=max(1, min(100, round(quality * 100))))
    return buffer.getvalue()


def save_copy(data: bytes, source: str | Path, extension: str, suffix: str = "") -> Path:
    source = Path(source)
    for index in range(10000):
        name = source.stem + suffix + (f" {index + 1}" if index else "") + "." + extension
        output = source.with_name(name)
        try:
            handle = output.open("xb")
        except FileExistsError:
            continue
        try:
            with handle:
                handle.write(data)
            return output
        except BaseException:
            output.unlink(missing_ok=True)
            raise
    raise ConversionError("Could not find an unused output filename.")


def render_page(page) -> Image.Image:
    width, height = page.get_size()
    if not all(math.isfinite(n) and n > 0 for n in (width, height)):
        raise ConversionError("This PDF page has invalid dimensions.")
    scale = min(300 / 72, 12000 / max(width, height))
    check_size(max(1, math.ceil(width * scale)), max(1, math.ceil(height * scale)))
    bitmap = page.render(scale=scale)
    try:
        return bitmap.to_pil().convert("RGBA")
    finally:
        bitmap.close()


def recognize_text(image: Image.Image) -> str:
    executable = shutil.which("tesseract")
    if not executable:
        raise ConversionError("Text recognition requires Tesseract. Install tesseract-ocr on Ubuntu or tesseract on Fedora/Windows, then try again.")
    with tempfile.TemporaryDirectory(prefix="shiftdrop-ocr-") as directory:
        path = Path(directory) / "scan.png"
        scan = image.copy()
        scan.thumbnail((5000, 5000))
        scan.convert("RGB").save(path)
        try:
            result = subprocess.run([executable, str(path), "stdout", "-l", "eng"], capture_output=True, timeout=90, check=False)
        except subprocess.TimeoutExpired as error:
            raise ConversionError("Text recognition timed out. Try a smaller image.") from error
        if result.returncode:
            raise ConversionError("Text recognition failed. Check that Tesseract's English language data is installed.")
        return result.stdout.decode("utf-8").strip()


def extract_text(source: Path, kind: str) -> str:
    if kind == "text":
        if source.suffix.lower() == ".docx":
            doc = Document(source)
            lines = [p.text for p in doc.paragraphs]
            lines += ["\t".join(cell.text for cell in row.cells) for table in doc.tables for row in table.rows]
            return "\n".join(lines)
        text = source.read_text(encoding="utf-8")
        return rtf_to_text(text) if source.suffix.lower() == ".rtf" else text
    if kind == "image":
        return recognize_text(load_image(source))
    pages = []
    with pdfium.PdfDocument(source) as document:
        for index in range(len(document)):
            with closing(document[index]) as page:
                with closing(page.get_textpage()) as textpage:
                    text = textpage.get_text_bounded().strip()
                pages.append(text or recognize_text(render_page(page)))
    return "\n\n".join(pages)


def text_pdf(text: str) -> bytes:
    output = io.BytesIO()
    styles = getSampleStyleSheet()
    story = []
    for line in text.splitlines() or [""]:
        story.append(Paragraph(escape(line) or "&#160;", styles["BodyText"]))
        story.append(Spacer(1, 3))
    SimpleDocTemplate(output, pagesize=(595, 842), leftMargin=48, rightMargin=48, topMargin=48, bottomMargin=48).build(story)
    return output.getvalue()


def images_pdf(images: list[Image.Image]) -> bytes:
    output = io.BytesIO()
    rgb = [image.convert("RGB") for image in images]
    rgb[0].save(output, format="PDF", save_all=True, append_images=rgb[1:], resolution=144)
    return output.getvalue()


def convert(source: str | Path, format_name: str) -> list[Path]:
    source = Path(source)
    kind = detect_kind(source)
    if format_name not in EXTENSIONS:
        raise ConversionError("Choose a supported output format.")
    if format_name in {"TXT", "DOCX"}:
        text = extract_text(source, kind)
        if not text.strip():
            raise ConversionError("No readable text was found in this file.")
        if format_name == "TXT":
            data = text.encode("utf-8")
        else:
            document = Document()
            for line in text.splitlines():
                document.add_paragraph(line)
            buffer = io.BytesIO()
            document.save(buffer)
            data = buffer.getvalue()
        return [save_copy(data, source, EXTENSIONS[format_name])]
    if kind == "text":
        pdf = text_pdf(extract_text(source, kind))
        if format_name == "PDF":
            return [save_copy(pdf, source, "pdf")]
        document = pdfium.PdfDocument(pdf)
    elif kind == "pdf":
        if format_name == "PDF":
            return [save_copy(source.read_bytes(), source, "pdf")]
        document = pdfium.PdfDocument(source)
    else:
        document = None
    outputs = []
    try:
        if document is not None:
            if not len(document):
                raise ConversionError("The PDF contains no pages.")
            for index in range(len(document)):
                with closing(document[index]) as page:
                    data = encode(render_page(page), format_name)
                suffix = f" Page {index + 1:03d}" if len(document) > 1 else ""
                outputs.append(save_copy(data, source, EXTENSIONS[format_name], suffix))
        else:
            if format_name == "PDF":
                return [save_copy(images_pdf([load_image(source)]), source, "pdf")]
            with Image.open(source) as original:
                count = original.n_frames if source.suffix.lower() in {".tif", ".tiff"} else 1
                for index in range(count):
                    original.seek(index)
                    check_size(*original.size)
                    image = ImageOps.exif_transpose(original).convert("RGBA")
                    suffix = f" Page {index + 1:03d}" if count > 1 else ""
                    outputs.append(save_copy(encode(image, format_name), source, EXTENSIONS[format_name], suffix))
        return outputs
    except BaseException:
        for output in outputs:
            output.unlink(missing_ok=True)
        raise
    finally:
        if document is not None:
            document.close()


def crop_image(image: Image.Image, box: tuple[float, float, float, float]) -> Image.Image:
    x, y, width, height = box
    if not all(math.isfinite(n) for n in box) or width <= 0 or height <= 0 or x < 0 or y < 0 or x + width > 1.000001 or y + height > 1.000001:
        raise ConversionError("Choose a crop area inside the image.")
    left, top = round(x * image.width), round(y * image.height)
    right, bottom = min(image.width, max(left + 1, round((x + width) * image.width))), min(image.height, max(top + 1, round((y + height) * image.height)))
    return image.crop((left, top, right, bottom))


PALETTES = [
    ("#ff457b", "#ff9d52", "#ffd437", "#6876ef"), ("#6b8ef3", "#ffffff"),
    ("#ffe35a", "#4edca8", "#df8edc"), ("#c990e4", "#f7f0ff"), ("#d09bdc", "#eebbbf"),
    ("#ee6a4c", "#439878", "#253d32"), ("#1769ef", "#32d8e8"), ("#00b4d4", "#93e9c5"),
    ("#ffffff", "#bcc8ef"), ("#ffffff",), ("#f7cc39",), ("#8ed534",), ("#4eac78",),
    ("#2d9e84",), ("#4bcbd0",), ("#446ce0",), ("#1d283c",), ("#793bb7",),
]


def add_background(image: Image.Image, palette=0, padding=160, corners=40, shadow=24, ratio=None, preview=False) -> Image.Image:
    from PIL import ImageColor
    padding = max(0, min(1000, int(padding)))
    width, height = image.width + padding * 2, image.height + padding * 2
    if ratio:
        if width / height > ratio:
            height = math.ceil(width / ratio)
        else:
            width = math.ceil(height * ratio)
    check_size(width, height)
    scale = min(1, 1000 / max(width, height)) if preview else 1
    width, height = max(1, round(width * scale)), max(1, round(height * scale))
    values = PALETTES[palette] if isinstance(palette, int) else (palette,)
    rgb = [ImageColor.getrgb(color) for color in values]
    if len(rgb) == 1:
        canvas = Image.new("RGBA", (width, height), rgb[0])
    else:
        # Render a small smooth color field, then scale it to the export resolution.
        field = Image.new("RGB", (160, 160))
        pixels = field.load()
        for y in range(160):
            for x in range(160):
                u, v = x / 159, y / 159
                if len(rgb) == 4:
                    weights = ((1-u)*(1-v), (1-u)*v, u*(1-v), u*v)
                    pixels[x, y] = tuple(round(sum(color[c] * weight for color, weight in zip(rgb, weights))) for c in range(3))
                else:
                    t = (u + v) / 2 * (len(rgb) - 1)
                    index = min(len(rgb) - 2, int(t)); fraction = t - index
                    pixels[x, y] = tuple(round(rgb[index][c] * (1-fraction) + rgb[index+1][c] * fraction) for c in range(3))
        canvas = field.resize((width, height), Image.Resampling.BICUBIC).convert("RGBA")
    size = (max(1, round(image.width * scale)), max(1, round(image.height * scale)))
    foreground = image.resize(size, Image.Resampling.LANCZOS).convert("RGBA")
    mask = Image.new("L", size)
    ImageDraw.Draw(mask).rounded_rectangle((0, 0, size[0]-1, size[1]-1), radius=max(0, corners * scale), fill=255)
    from PIL import ImageChops
    foreground.putalpha(ImageChops.multiply(mask, foreground.getchannel("A")))
    position = ((width-size[0])//2, (height-size[1])//2)
    if shadow > 0:
        shadow_mask = Image.new("L", canvas.size)
        shadow_mask.paste(mask, (position[0], position[1] + round(shadow * scale / 3)))
        shadow_mask = shadow_mask.filter(ImageFilter.GaussianBlur(shadow * scale)).point(lambda p: round(p * .35))
        shade = Image.new("RGBA", canvas.size, "black"); shade.putalpha(shadow_mask)
        canvas = Image.alpha_composite(canvas, shade)
    canvas.alpha_composite(foreground, position)
    return canvas


def adjust_image(image: Image.Image, brightness=1.0, contrast=1.0, saturation=1.0, rotation=0, flipped=False) -> Image.Image:
    alpha = image.convert("RGBA").getchannel("A")
    result = ImageEnhance.Brightness(image.convert("RGB")).enhance(brightness)
    result = ImageEnhance.Contrast(result).enhance(contrast)
    result = ImageEnhance.Color(result).enhance(saturation).convert("RGBA")
    result.putalpha(alpha)
    if flipped:
        result = ImageOps.mirror(result)
    return result.rotate(-90 * (rotation % 4), expand=True)
