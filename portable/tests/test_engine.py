import io
from contextlib import closing
from pathlib import Path

import pytest
from PIL import Image
import pypdfium2 as pdfium
from docx import Document

from shiftdrop import engine


@pytest.mark.parametrize('format_name', ['JPG', 'PNG', 'WEBP', 'HEIC', 'TIFF', 'AVIF', 'BMP'])
def test_real_image_codecs(image_file, format_name):
    output, = engine.convert(image_file, format_name)
    assert output != image_file
    with Image.open(output) as image:
        assert image.size == (120, 80)
        image.load()
    assert output.suffix == '.' + engine.EXTENSIONS[format_name]


def test_source_and_colliding_files_are_preserved(image_file):
    original = image_file.read_bytes()
    occupied = image_file.with_suffix('.jpg'); occupied.write_bytes(b'existing file')
    first, = engine.convert(image_file, 'JPG')
    second, = engine.convert(image_file, 'JPG')
    assert first.name == 'Original 2.jpg'
    assert second.name == 'Original 3.jpg'
    assert occupied.read_bytes() == b'existing file'
    assert image_file.read_bytes() == original


def test_pdf_image_and_text_roundtrip(tmp_path):
    source = tmp_path / 'Notes.txt'; source.write_text('A & B < C > D\nThis stays local.', encoding='utf-8')
    pdf, = engine.convert(source, 'PDF')
    with pdfium.PdfDocument(pdf) as document:
        assert len(document) == 1
        with closing(document[0]) as page, closing(page.get_textpage()) as text:
            assert 'A & B < C > D' in text.get_text_bounded()
    image, = engine.convert(pdf, 'PNG')
    with Image.open(image) as decoded:
        assert decoded.width > 2000
    word, = engine.convert(pdf, 'DOCX')
    assert 'A & B < C > D' in '\n'.join(p.text for p in Document(word).paragraphs)
    text, = engine.convert(pdf, 'TXT')
    assert 'This stays local.' in text.read_text(encoding='utf-8')


def test_multipage_pdf_exports_every_page(tmp_path):
    source = tmp_path / 'Pages.pdf'
    from reportlab.pdfgen.canvas import Canvas
    canvas = Canvas(str(source), pagesize=(120, 80)); canvas.drawString(10, 40, 'One'); canvas.showPage()
    canvas.setPageSize((80, 120)); canvas.drawString(10, 40, 'Two'); canvas.showPage(); canvas.save()
    outputs = engine.convert(source, 'PNG')
    assert [path.name for path in outputs] == ['Pages Page 001.png', 'Pages Page 002.png']
    with Image.open(outputs[0]) as first, Image.open(outputs[1]) as second:
        assert first.width > first.height
        assert second.height > second.width


def test_pdf_batch_rolls_back_only_its_own_outputs(tmp_path, monkeypatch):
    source = tmp_path / 'Pages.pdf'
    from reportlab.pdfgen.canvas import Canvas
    canvas = Canvas(str(source), pagesize=(120, 80))
    for number in range(2): canvas.drawString(10, 40, str(number)); canvas.showPage()
    canvas.save()
    existing = tmp_path / 'Pages Page 001.png'; existing.write_bytes(b'preserve')
    original = engine.encode; calls = []
    def encode(image, format_name, quality=.9):
        calls.append(1)
        if len(calls) == 2: raise OSError('disk full')
        return original(image, format_name, quality)
    monkeypatch.setattr(engine, 'encode', encode)
    with pytest.raises(OSError, match='disk full'): engine.convert(source, 'PNG')
    assert sorted(path.name for path in tmp_path.iterdir()) == ['Pages Page 001.png', 'Pages.pdf']
    assert existing.read_bytes() == b'preserve'


def test_long_text_is_paginated(tmp_path):
    source = tmp_path / 'Long.txt'; source.write_text('\n'.join(f'Line {n}: Local conversion.' for n in range(120)), encoding='utf-8')
    output, = engine.convert(source, 'PDF')
    with pdfium.PdfDocument(output) as document:
        assert len(document) > 1
        with closing(document[len(document)-1]) as page, closing(page.get_textpage()) as text:
            assert 'Line 119:' in text.get_text_bounded()


def test_invalid_input_and_pixel_limits(tmp_path):
    with pytest.raises(engine.ConversionError): engine.detect_kind(tmp_path)
    source = tmp_path / 'Broken.png'; source.write_bytes(b'not an image')
    with pytest.raises(Exception): engine.convert(source, 'JPG')
    assert list(tmp_path.iterdir()) == [source]
    with pytest.raises(engine.ConversionError): engine.check_size(16000, 16000)
    with pytest.raises(engine.ConversionError): engine.check_size(0, 80)


def test_missing_ocr_is_actionable(image_file, monkeypatch):
    monkeypatch.setattr(engine.shutil, 'which', lambda _: None)
    with pytest.raises(engine.ConversionError, match='Tesseract'):
        engine.convert(image_file, 'DOCX')


def test_crop_background_adjustments_and_metadata_removal(image_file):
    image = engine.load_image(image_file)
    cropped = engine.crop_image(image, (0, 0, .5, .5))
    assert cropped.size == (60, 40)
    assert cropped.getpixel((10, 10))[:3] == (25, 100, 220)
    output = engine.add_background(cropped, padding=20, corners=10, shadow=4, ratio=1)
    assert output.width == output.height == 100
    rotated = engine.adjust_image(cropped, rotation=1)
    assert rotated.size == (40, 60)
    image.info['exif'] = Image.Exif().tobytes()
    image.info['comment'] = b'private metadata'
    with Image.open(io.BytesIO(engine.encode(image, 'PNG'))) as result:
        assert not result.getexif()
        assert 'comment' not in result.info
    with pytest.raises(engine.ConversionError): engine.crop_image(image, (.9, 0, .5, 1))
